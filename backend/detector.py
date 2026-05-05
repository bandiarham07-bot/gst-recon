import openpyxl
from pathlib import Path


def detect_file_type(path: str) -> str:
    wb = openpyxl.load_workbook(path, read_only=True)
    result = "2B" if "Read me" in wb.sheetnames else "EPR"
    wb.close()
    return result


def get_sheet_names(path: str):
    wb = openpyxl.load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def peek_rows(path: str, sheet_name: str, n: int = 20):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet_name]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= n:
            break
        rows.append(row)
    wb.close()
    return rows
