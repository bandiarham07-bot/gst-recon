import uuid
import openpyxl
from typing import List, Dict, Tuple, Optional

from backend.ingestion.canonical_map import map_column, normalise, _load_registry

try:
    from rapidfuzz import fuzz
    HAS_RAPIDFUZZ = True
except ImportError:
    HAS_RAPIDFUZZ = False


def _is_numeric(val) -> bool:
    if val is None:
        return False
    if isinstance(val, (int, float)):
        return True
    try:
        float(str(val).replace(",", ""))
        return True
    except ValueError:
        return False


def _is_date(val) -> bool:
    import datetime
    return isinstance(val, (datetime.date, datetime.datetime))


def _alias_match_score(cell_val: str) -> float:
    if not cell_val or not cell_val.strip():
        return 0.0
    registry = _load_registry()
    norm = normalise(str(cell_val))
    best = 0.0
    for aliases in registry.values():
        for alias in aliases:
            if HAS_RAPIDFUZZ:
                score = fuzz.token_sort_ratio(norm, normalise(alias)) / 100.0
            else:
                score = 1.0 if norm == normalise(alias) else 0.0
            if score > best:
                best = score
    return best


def score_row(row: tuple, next_rows: List[tuple]) -> float:
    cells = [c for c in row if c is not None]
    if not cells:
        return 0.0

    total = len(row) if row else 1

    string_cells = [c for c in cells if isinstance(c, str)]
    numeric_cells = [c for c in cells if _is_numeric(c) and not isinstance(c, str)]
    date_cells = [c for c in cells if _is_date(c)]

    string_density = len(string_cells) / max(len(cells), 1)
    numeric_absence = 1.0 - (len(numeric_cells) + len(date_cells)) / max(len(cells), 1)

    alias_scores = [_alias_match_score(str(c)) for c in string_cells]
    alias_match = sum(1 for s in alias_scores if s > 0.8) / max(len(string_cells), 1)

    unique_vals = set(str(c).lower().strip() for c in cells if c is not None)
    uniqueness = len(unique_vals) / max(len(cells), 1)

    density_jump = 0.0
    for next_row in next_rows[:3]:
        next_cells = [c for c in next_row if c is not None]
        if not next_cells:
            continue
        next_numeric = [c for c in next_cells if _is_numeric(c) and not isinstance(c, str)]
        next_density = len(next_numeric) / max(len(next_cells), 1)
        if next_density > 0.3:
            density_jump = 0.3
            break

    score = (
        string_density * 0.25 +
        numeric_absence * 0.20 +
        alias_match * 0.30 +
        uniqueness * 0.15 +
        density_jump * 0.10
    )
    return min(score, 1.0)


def detect_header_row(path: str, sheet_name: str) -> Tuple[Optional[int], float, List[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    max_scan = min(20, len(rows))
    candidates = []

    for i in range(max_scan):
        next_rows = rows[i + 1:i + 4]
        score = score_row(rows[i], next_rows)
        if score > 0.3:
            candidates.append({"row_index": i, "score": score, "row": rows[i]})

    if not candidates:
        return None, 0.0, []

    candidates.sort(key=lambda x: x["score"], reverse=True)
    best = candidates[0]
    return best["row_index"], best["score"], candidates[:3]


def _safe_float(val) -> Optional[float]:
    try:
        if val is None or val == "":
            return None
        return float(str(val).replace(",", ""))
    except (ValueError, TypeError):
        return None


def _safe_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def parse_epr(
    path: str,
    batch_id: str,
    epr_software: str,
    header_row_idx: int,
    col_mapping: Dict[str, str]
) -> List[dict]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_row = rows[header_row_idx]
    raw_to_idx: Dict[str, int] = {}
    for i, cell in enumerate(header_row):
        if cell is not None:
            raw_to_idx[str(cell).strip()] = i

    canonical_to_idx: Dict[str, int] = {}
    for raw, canonical in col_mapping.items():
        if raw in raw_to_idx and canonical:
            canonical_to_idx[canonical] = raw_to_idx[raw]

    monetary_fields = {"invoice_value", "taxable_value", "igst", "cgst", "sgst", "cess"}
    parsed_rows = []

    for row in rows[header_row_idx + 1:]:
        if not any(row):
            continue
        record = {
            "row_id": str(uuid.uuid4()),
            "batch_id": batch_id,
            "epr_software": epr_software,
            "source_type": "PR",
            "recon_status": None,
        }
        for canonical, idx in canonical_to_idx.items():
            if idx < len(row):
                val = row[idx]
                if canonical in monetary_fields:
                    record[canonical] = _safe_float(val)
                else:
                    record[canonical] = _safe_str(val)
        parsed_rows.append(record)

    return parsed_rows
