import uuid
import openpyxl
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple, Optional

from backend.database.schema import SHEET_TABLE_MAP, SUMMARY_SHEETS
from backend.ingestion.canonical_map import map_column, normalise

DATA_SHEETS = [s for s in SHEET_TABLE_MAP.keys()]

B2B_HEADER_ROWS = (4, 5)  # 0-indexed rows 4 and 5


def _parse_readme(path: str) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Read me"]
    meta = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] == "Financial Year":
            meta["financial_year"] = str(row[2]) if row[2] else ""
        elif row[0] == "Tax Period":
            meta["tax_period"] = str(row[2]) if row[2] else ""
        elif row[0] == "GSTIN":
            meta["company_gstin"] = str(row[2]) if row[2] else ""
        elif row[0] == "Legal Name":
            meta["legal_name"] = str(row[2]) if row[2] else ""
        elif row[0] == "Date of generation":
            meta["date_of_generation"] = str(row[2]) if row[2] else ""
    wb.close()
    return meta


def _merge_header_rows(row4, row5) -> List[str]:
    merged = []
    for i, (parent, child) in enumerate(zip(row4, row5)):
        parent_str = str(parent).strip() if parent else ""
        child_str = str(child).strip() if child else ""
        if parent_str and child_str:
            merged.append(f"{parent_str}__{child_str}")
        elif child_str:
            merged.append(child_str)
        elif parent_str:
            merged.append(parent_str)
        else:
            merged.append("")
    return merged


def _map_merged_headers(merged_headers: List[str]) -> Dict[int, str]:
    col_map = {}
    for i, raw in enumerate(merged_headers):
        if not raw:
            continue
        canonical, score = map_column(raw)
        if not canonical:
            canonical, score = map_column(raw.split("__")[-1])
        if canonical:
            col_map[i] = canonical
    return col_map


def _safe_float(val) -> Optional[float]:
    try:
        if val is None or val == "":
            return None
        return float(val)
    except (ValueError, TypeError):
        return None


def _safe_str(val) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _parse_data_sheet(ws, sheet_name: str, batch_id: str) -> List[dict]:
    rows_data = list(ws.iter_rows(values_only=True))
    if len(rows_data) < 6:
        return []

    row4 = rows_data[4] if len(rows_data) > 4 else []
    row5 = rows_data[5] if len(rows_data) > 5 else []

    merged_headers = _merge_header_rows(
        row4 + (None,) * max(0, len(row5) - len(row4)),
        row5 + (None,) * max(0, len(row4) - len(row5))
    )
    col_map = _map_merged_headers(merged_headers)

    parsed_rows = []
    table = SHEET_TABLE_MAP.get(sheet_name, "gstr2b_b2b")

    for row in rows_data[6:]:
        if not any(row):
            continue
        if row[0] is None or str(row[0]).strip() == "":
            continue

        record = {
            "row_id": str(uuid.uuid4()),
            "batch_id": batch_id,
            "sheet_type": sheet_name,
            "source_type": "2B",
            "recon_status": None,
        }

        for col_idx, canonical in col_map.items():
            if col_idx < len(row):
                val = row[col_idx]
                monetary_fields = {"invoice_value", "taxable_value", "igst", "cgst", "sgst", "cess"}
                if canonical in monetary_fields:
                    record[canonical] = _safe_float(val)
                else:
                    record[canonical] = _safe_str(val)

        parsed_rows.append(record)

    return parsed_rows


def _parse_summary_sheet(ws, sheet_name: str, batch_id: str) -> List[dict]:
    rows_data = list(ws.iter_rows(values_only=True))
    summaries = []
    header_found = False
    for row in rows_data:
        if not any(row):
            continue
        if row[0] == "S.no.":
            header_found = True
            continue
        if header_found and row[0] is not None and str(row[0]).strip().isdigit() is False:
            if isinstance(row[0], str) and row[0].startswith(("I", "II", "III", "IV", "V", "Part")):
                continue
        if header_found and row[1] and str(row[1]).strip():
            try:
                igst = _safe_float(row[3]) or 0.0
                cgst = _safe_float(row[4]) or 0.0
                sgst = _safe_float(row[5]) or 0.0
                cess = _safe_float(row[6]) or 0.0
                summaries.append({
                    "batch_id": batch_id,
                    "summary_type": sheet_name,
                    "heading": _safe_str(row[1]),
                    "gstr3b_table": _safe_str(row[2]) if len(row) > 2 else "",
                    "igst": igst,
                    "cgst": cgst,
                    "sgst": sgst,
                    "cess": cess,
                })
            except Exception:
                pass
    return summaries


def parse_2b(path: str, batch_id: str) -> Tuple[dict, Dict[str, List[dict]], List[dict]]:
    meta = _parse_readme(path)
    wb = openpyxl.load_workbook(path, read_only=True)

    all_data: Dict[str, List[dict]] = {}
    itc_summaries: List[dict] = []

    for sheet_name in wb.sheetnames:
        if sheet_name == "Read me":
            continue
        ws = wb[sheet_name]
        if sheet_name in SUMMARY_SHEETS:
            itc_summaries.extend(_parse_summary_sheet(ws, sheet_name, batch_id))
        elif sheet_name in SHEET_TABLE_MAP:
            table = SHEET_TABLE_MAP[sheet_name]
            rows = _parse_data_sheet(ws, sheet_name, batch_id)
            if rows:
                if table not in all_data:
                    all_data[table] = []
                all_data[table].extend(rows)

    wb.close()
    return meta, all_data, itc_summaries
